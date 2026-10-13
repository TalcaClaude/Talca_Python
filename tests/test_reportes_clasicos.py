"""
Tests de los reportes clásicos (los mismos que hoy saca Talca PHP), empezando
por el "Reporte de OT" de Operativo (/api/operativo/reportes/ordenes).

Criterio copiado de createXLS_OT de Talca PHP:
  - estado=abiertas  -> OT no cerradas, filtradas por fecha de creación
  - estado=cerradas  -> OT entregadas/canceladas, filtradas por fecha de cierre
                        (fecha_cierre_talca si vino del sync; si no, la de entrega)
  - estado=todas     -> todas, filtradas por fecha de creación
  - las OT borradas en Talca nunca salen
Montos: el total de la OT es el precio (+ IVA), NO la suma de los ítems; los
gastos son repuestos + mano de obra con costo + convertidores + caja chica por
OT; margen = neto - gastos.

Cómo correrlos:
    python tests/correr_todo.py reportes_clasicos
"""
import io
import os
import sys
import tempfile
import unittest

_TMP_DB_FD, _TMP_DB_PATH = tempfile.mkstemp(prefix="test_taller_reportes_", suffix=".db")
os.close(_TMP_DB_FD)
os.remove(_TMP_DB_PATH)
os.environ["DB_PATH"] = _TMP_DB_PATH
os.environ.setdefault("SECRET_KEY", "clave-de-pruebas")

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import app as taller_app  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

TS = "2026-10-13T10:00:00"


class TestReporteOT(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        taller_app.app.config["TESTING"] = True
        cls.client = taller_app.app.test_client()
        r = cls.client.post("/login", data={"username": "admin", "password": "taller2026"})
        assert r.status_code in (200, 302), r.status_code

        db = taller_app._conectar_db()
        cli = db.execute("INSERT INTO clientes (nombre, creado_en) VALUES ('CLIENTE REPORTE', ?)", (TS,)).lastrowid
        eq = db.execute(
            "INSERT INTO equipos (cliente_id, identificador, marca, modelo, color, anio, creado_en) "
            "VALUES (?, 'ABCD12', 'TOYOTA', 'RAV4', 'ROJO', 2020, ?)", (cli, TS)).lastrowid

        def orden(estado, recepcion, **extra):
            campos = {"cliente_id": cli, "equipo_id": eq, "estado": estado, "fecha_recepcion": recepcion,
                      "creado_en": TS, "actualizado_en": TS, **extra}
            cols = ", ".join(campos)
            marcas = ", ".join("?" for _ in campos)
            return db.execute(f"INSERT INTO ordenes_trabajo ({cols}) VALUES ({marcas})", list(campos.values())).lastrowid

        # A: abierta, creada en octubre. Precio 100.000 + IVA = 119.000; pagó 50.000;
        # un repuesto de 20.000 y una mano de obra sin costo.
        cls.a = orden("recibido", "2026-10-05T09:00:00", precio_referencial=100000, precio_iva=1,
                      trabajo_a_realizar="OVERHAUL COMPLETO", referencia="GOOGLE")
        db.execute("INSERT INTO orden_items (orden_id, tipo, descripcion, cantidad, costo_unitario, creado_en) "
                   "VALUES (?, 'repuesto', 'KIT', 1, 20000, ?)", (cls.a, TS))
        db.execute("INSERT INTO orden_items (orden_id, tipo, descripcion, cantidad, costo_unitario, creado_en) "
                   "VALUES (?, 'mano_obra', 'MO', 1, 0, ?)", (cls.a, TS))
        db.execute("INSERT INTO pagos_orden (orden_id, metodo_pago, monto, fecha, creado_en) "
                   "VALUES (?, 'efectivo', 50000, '2026-10-06', ?)", (cls.a, TS))
        # B: creada en septiembre, entregada (cerrada) en octubre, sin IVA.
        cls.b = orden("entregado", "2026-09-20T09:00:00", fecha_entrega_real="2026-10-02T17:00:00",
                      precio_referencial=50000)
        # C: borrada en Talca: no sale nunca.
        cls.c = orden("recibido", "2026-10-07T09:00:00", borrada_en_talca=1, precio_referencial=999)
        # D: vino del sync (talca_id), creada en agosto, cerrada en Talca en octubre.
        cls.d = orden("entregado", "2026-08-01T09:00:00", talca_id=8100, estado_talca="CERRADA BLOQUEADA",
                      fecha_cierre_talca="2026-10-03 12:00:00", fecha_entrega_real="2026-10-04T12:00:00",
                      tipo_talca="Nueva Orden", patente="ZZ9999", precio_referencial=30000)
        db.commit()
        db.close()

    def _filas(self, **params):
        resp = self.client.get("/api/operativo/reportes/ordenes", query_string=params)
        if resp.status_code != 200:  # el cuerpo es un Excel: solo se lee como texto si falló
            self.fail(f"HTTP {resp.status_code}: {resp.data[:300]!r}")
        ws = load_workbook(io.BytesIO(resp.data)).active
        filas = list(ws.iter_rows(values_only=True))
        i = next(k for k, f in enumerate(filas) if f and f[0] == "N° OT")
        enc = filas[i]
        datos = [dict(zip(enc, f)) for f in filas[i + 1:] if f and f[0] not in (None, "")]
        return datos

    def _numeros(self, **params):
        return sorted(str(f["N° OT"]) for f in self._filas(**params))

    def test_abiertas_por_fecha_de_creacion(self):
        self.assertEqual([f"L-{self.a}"], self._numeros(estado="abiertas", desde="2026-10-01", hasta="2026-10-31"))

    def test_cerradas_por_fecha_de_cierre(self):
        # B y D se cerraron en octubre aunque se crearon antes; A no está cerrada.
        self.assertEqual(sorted([f"L-{self.b}", "8100"]),
                         self._numeros(estado="cerradas", desde="2026-10-01", hasta="2026-10-31"))
        self.assertEqual([], self._numeros(estado="cerradas", desde="2026-09-01", hasta="2026-09-30"))

    def test_cierre_del_sync_manda_sobre_la_entrega(self):
        # D se cerró en Talca el 03-10 y se entregó el 04-10: el reporte usa el cierre.
        self.assertEqual(["8100"], self._numeros(estado="cerradas", desde="2026-10-03", hasta="2026-10-03"))

    def test_todas_por_fecha_de_creacion(self):
        self.assertEqual([f"L-{self.a}"], self._numeros(estado="todas", desde="2026-10-01", hasta="2026-10-31"))
        self.assertEqual([f"L-{self.b}"], self._numeros(estado="todas", desde="2026-09-01", hasta="2026-09-30"))

    def test_las_borradas_en_talca_no_salen(self):
        for estado in ("todas", "abiertas", "cerradas"):
            with self.subTest(estado=estado):
                self.assertNotIn(f"L-{self.c}", self._numeros(estado=estado))

    def test_montos_salen_del_precio_y_no_de_los_items(self):
        a = next(f for f in self._filas(estado="abiertas") if f["N° OT"] == f"L-{self.a}")
        self.assertEqual(119000, a["Total OT"])       # 100.000 + IVA
        self.assertEqual(50000, a["Total cobrado"])
        self.assertEqual(69000, a["Diferencia"])
        self.assertEqual(20000, a["Total gastos"])
        self.assertEqual(80000, a["Margen"])          # neto 100.000 - gastos 20.000
        self.assertEqual("SI", a["Con IVA"])
        self.assertEqual("OVERHAUL COMPLETO", a["Requerimiento"])
        self.assertEqual("ABCD12", a["Patente"])

    def test_la_patente_de_la_ot_manda_sobre_la_del_vehiculo(self):
        d = next(f for f in self._filas(estado="cerradas") if str(f["N° OT"]) == "8100")
        self.assertEqual("ZZ9999", d["Patente"])
        self.assertEqual("CERRADA BLOQUEADA", d["Estado en Talca"])
        self.assertEqual("2026-10-03", d["F. cierre"])
        self.assertEqual(63, d["Días a cierre"])

    def test_estado_desconocido_cae_a_todas(self):
        self.assertEqual(self._numeros(estado="todas"), self._numeros(estado="cualquiera"))

    def test_pdf_responde(self):
        resp = self.client.get("/api/operativo/reportes/ordenes/pdf", query_string={"estado": "cerradas"})
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(resp.data.startswith(b"%PDF"))

    def test_resumen_cuenta_con_el_mismo_criterio(self):
        r = self.client.get("/api/operativo/reportes/resumen",
                            query_string={"ordenes_estado": "abiertas", "ordenes_desde": "2026-10-01"}).get_json()
        self.assertEqual(1, r["ordenes"]["total"])
        self.assertEqual(119000, r["ordenes"]["total_facturado"])


class TestReporteGastosOT(unittest.TestCase):
    """Reporte de gastos por OT (/api/operativo/reportes/gastos-ot): como
    createXLS_sell_nuevo de Talca PHP, una fila por gasto de cada OT."""

    @classmethod
    def setUpClass(cls):
        taller_app.app.config["TESTING"] = True
        cls.client = taller_app.app.test_client()
        cls.client.post("/login", data={"username": "admin", "password": "taller2026"})
        db = taller_app._conectar_db()
        cli = db.execute("INSERT INTO clientes (nombre, creado_en) VALUES ('CLIENTE GASTOS', ?)", (TS,)).lastrowid
        cls.ot = db.execute(
            "INSERT INTO ordenes_trabajo (cliente_id, estado, fecha_recepcion, creado_en, actualizado_en, talca_id) "
            "VALUES (?, 'recibido', '2025-03-10T09:00:00', ?, ?, 8300)", (cli, TS, TS)).lastrowid
        items = [("repuesto", "DISCOS", 2, 15000, 0), ("mano_obra", "TORNERIA EXTERNA", 1, 10000, 1),
                 ("mano_obra", "MO INTERNA", 1, 0, 0)]
        for tipo, desc, cant, costo, iva in items:
            db.execute("INSERT INTO orden_items (orden_id, tipo, descripcion, cantidad, costo_unitario, aplica_iva, "
                       "creado_en) VALUES (?, ?, ?, ?, ?, ?, '2025-03-11T10:00:00')", (cls.ot, tipo, desc, cant, costo, iva))
        gasto = db.execute("INSERT INTO caja_movimientos (tipo, medio_pago, fecha, monto, creado_en) "
                           "VALUES ('gasto', 'efectivo_cajero', '2025-03-12T10:00:00', 8000, ?)", (TS,)).lastrowid
        db.execute("INSERT INTO caja_gasto_detalles (gasto_id, numero_documento, monto, rubro_gav, descripcion, "
                   "orden_id, creado_en) VALUES (?, 'B-77', 8000, 'gastos_por_ot', 'RETEN', ?, ?)", (gasto, cls.ot, TS))
        db.commit()
        db.close()

    def _filas(self, **params):
        resp = self.client.get("/api/operativo/reportes/gastos-ot", query_string=params)
        if resp.status_code != 200:
            self.fail(f"HTTP {resp.status_code}: {resp.data[:300]!r}")
        filas = list(load_workbook(io.BytesIO(resp.data)).active.iter_rows(values_only=True))
        i = next(k for k, f in enumerate(filas) if f and f[0] == "N° OT")
        return [dict(zip(filas[i], f)) for f in filas[i + 1:] if f and f[0] not in (None, "")]

    def test_una_fila_por_gasto_con_costo(self):
        filas = self._filas(estado="abiertas", desde="2025-03-01", hasta="2025-03-31")
        self.assertEqual(["Repuesto", "Mano de obra / externo", "Caja chica"], [f["Tipo de gasto"] for f in filas])
        self.assertEqual([30000, 11900, 8000], [f["Costo"] for f in filas])  # la externa lleva IVA
        self.assertTrue(all(str(f["N° OT"]) == "8300" for f in filas))

    def test_la_mano_de_obra_interna_sin_costo_no_sale(self):
        filas = self._filas(estado="abiertas", desde="2025-03-01", hasta="2025-03-31")
        self.assertNotIn("MO INTERNA", [f["Descripción"] for f in filas])

    def test_la_boleta_de_caja_chica_trae_su_documento_y_fecha(self):
        caja = next(f for f in self._filas(estado="abiertas", desde="2025-03-01", hasta="2025-03-31")
                    if f["Tipo de gasto"] == "Caja chica")
        self.assertEqual("B-77", caja["Documento"])
        self.assertEqual("2025-03-12", caja["F. gasto"])

    def test_fuera_del_rango_no_trae_nada(self):
        self.assertEqual([], self._filas(estado="abiertas", desde="2025-04-01", hasta="2025-04-30"))


class TestReporteCompras(unittest.TestCase):
    """Reporte de compras (/api/reportes/compras): como createXLS_rc de Talca
    PHP, un documento por fila con RUT, condición de pago y neto/IVA/total."""

    @classmethod
    def setUpClass(cls):
        taller_app.app.config["TESTING"] = True
        cls.client = taller_app.app.test_client()
        cls.client.post("/login", data={"username": "admin", "password": "taller2026"})
        db = taller_app._conectar_db()
        prov = db.execute("INSERT INTO proveedores (nombre, rut, creado_en) VALUES ('REPUESTOS SUR', '76.111.222-3', ?)",
                          (TS,)).lastrowid
        prod = db.execute(
            "INSERT INTO productos (codigo_interno, nombre, descripcion, categoria, marca, modelo, ubicacion, "
            "stock_disponible, stock_reservado, stock_merma, stock_minimo, costo_compra, creado_en, actualizado_en) "
            "VALUES ('', 'X', '', 'A', '', 'B', '', 0, 0, 0, 0, 0, ?, ?)", (TS, TS)).lastrowid
        # Factura nueva: trae sus montos (Migración 33).
        nueva = db.execute(
            "INSERT INTO facturas (numero, proveedor_id, fecha, creado_en, condicion_pago, monto_neto, monto_iva, "
            "monto_total, tipo_documento) VALUES ('F-1', ?, '2025-05-10', ?, 'credito', 100000, 19000, 119000, 'factura')",
            (prov, TS)).lastrowid
        db.execute("INSERT INTO factura_lineas (factura_id, producto_id, cantidad, precio_unitario) VALUES (?, ?, 2, 50000)",
                   (nueva, prod))
        # Factura antigua sin montos guardados: se calcula desde las líneas.
        vieja = db.execute(
            "INSERT INTO facturas (numero, proveedor, fecha, creado_en) VALUES ('F-0', 'PROVEEDOR SUELTO', '2025-05-02', ?)",
            (TS,)).lastrowid
        db.execute("INSERT INTO factura_lineas (factura_id, producto_id, cantidad, precio_unitario) VALUES (?, ?, 1, 10000)",
                   (vieja, prod))
        db.commit()
        db.close()

    def _filas(self, **params):
        resp = self.client.get("/api/reportes/compras", query_string=params)
        if resp.status_code != 200:
            self.fail(f"HTTP {resp.status_code}: {resp.data[:300]!r}")
        filas = list(load_workbook(io.BytesIO(resp.data)).active.iter_rows(values_only=True))
        i = next(k for k, f in enumerate(filas) if f and f[0] == "Fecha")
        return [dict(zip(filas[i], f)) for f in filas[i + 1:] if f and f[0] not in (None, "")]

    def test_columnas_de_talca_php(self):
        f = next(x for x in self._filas(desde="2025-05-10", hasta="2025-05-10"))
        self.assertEqual(("F-1", "REPUESTOS SUR", "76.111.222-3", "credito"),
                         (f["N° documento"], f["Proveedor"], f["RUT"], f["Condición de pago"]))
        self.assertEqual((100000, 19000, 119000), (f["Neto"], f["IVA"], f["Total"]))

    def test_factura_sin_montos_se_calcula_desde_las_lineas(self):
        f = next(x for x in self._filas(desde="2025-05-02", hasta="2025-05-02"))
        self.assertEqual((10000, 1900, 11900), (f["Neto"], f["IVA"], f["Total"]))
        self.assertEqual("PROVEEDOR SUELTO", f["Proveedor"])

    def test_filtra_por_fecha_del_documento(self):
        self.assertEqual(["F-1", "F-0"], [f["N° documento"] for f in self._filas(desde="2025-05-01", hasta="2025-05-31")])

    def test_pdf_responde(self):
        resp = self.client.get("/api/reportes/compras/pdf", query_string={"desde": "2025-05-01"})
        self.assertEqual(200, resp.status_code)
        self.assertTrue(resp.data.startswith(b"%PDF"))


class TestReporteOTConvertidores(unittest.TestCase):
    """Reporte de OT de convertidores (/api/convertidores/reportes/ot): como
    createXLS_ot_convertidor de Talca PHP, más la hoja de costos por mes."""

    @classmethod
    def setUpClass(cls):
        taller_app.app.config["TESTING"] = True
        cls.client = taller_app.app.test_client()
        cls.client.post("/login", data={"username": "admin", "password": "taller2026"})
        db = taller_app._conectar_db()
        cli = db.execute("INSERT INTO clientes (nombre, rut, creado_en) VALUES ('CLIENTE CONV', '11.222.333-4', ?)",
                         (TS,)).lastrowid
        ot = db.execute("INSERT INTO ordenes_trabajo (cliente_id, estado, fecha_recepcion, creado_en, actualizado_en, "
                        "talca_id) VALUES (?, 'recibido', '2024-01-01T09:00:00', ?, ?, 8400)", (cli, TS, TS)).lastrowid
        # Externo: precio 200.000, reparación 50.000, cerrado en julio.
        cls.externo = db.execute(
            "INSERT INTO convertidores_torque (origen_tipo, cliente_id, marca, modelo, numero_serie, estado, "
            "precio_venta, fecha_recepcion, fecha_cierre) VALUES ('cliente_directo', ?, 'TOYOTA', 'U660', 'S1', "
            "'entregado', 200000, '2024-06-20T09:00:00', '2024-07-15T10:00:00')", (cli,)).lastrowid
        db.execute("INSERT INTO convertidor_items (convertidor_id, tipo, descripcion, cantidad, costo_unitario, creado_en) "
                   "VALUES (?, 'repuesto', 'KIT', 1, 50000, ?)", (cls.externo, TS))
        # Interno (desde la caja de un auto): abierto, sin precio, reparación 30.000.
        cls.interno = db.execute(
            "INSERT INTO convertidores_torque (origen_tipo, cliente_id, orden_origen_id, estado, precio_venta, "
            "fecha_recepcion) VALUES ('desde_ot_auto', ?, ?, 'reparacion', 99999, '2024-07-01T09:00:00')",
            (cli, ot)).lastrowid
        db.execute("INSERT INTO convertidor_items (convertidor_id, tipo, descripcion, cantidad, costo_unitario, creado_en) "
                   "VALUES (?, 'repuesto', 'DISCOS', 1, 30000, ?)", (cls.interno, TS))
        db.commit()
        db.close()

    def _libro(self, **params):
        resp = self.client.get("/api/convertidores/reportes/ot", query_string=params)
        if resp.status_code != 200:
            self.fail(f"HTTP {resp.status_code}: {resp.data[:300]!r}")
        return load_workbook(io.BytesIO(resp.data))

    def _filas(self, **params):
        filas = list(self._libro(**params)["OT de convertidores"].iter_rows(values_only=True))
        i = next(k for k, f in enumerate(filas) if f and f[0] == "N°")
        return [dict(zip(filas[i], f)) for f in filas[i + 1:] if f and isinstance(f[0], int)]

    def test_cerradas_por_fecha_de_cierre_con_utilidad_real(self):
        filas = self._filas(estado="cerradas", desde="2024-07-01", hasta="2024-07-31")
        self.assertEqual([self.externo], [f["N°"] for f in filas])
        f = filas[0]
        self.assertEqual((200000, 50000, 150000), (f["Precio de venta"], f["Costo reparación"], f["Utilidad"]))
        self.assertEqual("11.222.333-4", f["RUT"])

    def test_abiertas_por_recepcion_y_el_interno_no_tiene_precio(self):
        filas = self._filas(estado="abiertas", desde="2024-07-01", hasta="2024-07-31")
        self.assertEqual([self.interno], [f["N°"] for f in filas])
        f = filas[0]
        self.assertEqual("", f["Precio de venta"] or "")   # interno: su precio_venta no se usa
        self.assertEqual(30000, f["Costo reparación"])
        self.assertEqual("8400", str(f["OT de origen"]))

    def test_costos_por_mes(self):
        hoja = self._libro(estado="todas")["Costos por mes"]
        filas = [r for r in hoja.iter_rows(min_row=2, values_only=True) if r[0]]
        julio = [r for r in filas if r[0] == "2024-07"]
        self.assertEqual(1, len(julio))
        self.assertEqual((1, 50000, 200000), (julio[0][2], julio[0][3], julio[0][4]))


class TestReporteCobros(unittest.TestCase):
    """Reporte de cobros (/api/reportes/cobros): como createXLS_pay de Talca PHP,
    pagos de OT filtrados por fecha de pago, con el N° de OT de Talca."""

    @classmethod
    def setUpClass(cls):
        taller_app.app.config["TESTING"] = True
        cls.client = taller_app.app.test_client()
        cls.client.post("/login", data={"username": "admin", "password": "taller2026"})
        db = taller_app._conectar_db()
        cli = db.execute("INSERT INTO clientes (nombre, creado_en) VALUES ('CLIENTE COBROS', ?)", (TS,)).lastrowid
        cls.local = db.execute(
            "INSERT INTO ordenes_trabajo (cliente_id, estado, fecha_recepcion, creado_en, actualizado_en) "
            "VALUES (?, 'recibido', '2025-01-15T09:00:00', ?, ?)", (cli, TS, TS)).lastrowid
        cls.talca = db.execute(
            "INSERT INTO ordenes_trabajo (cliente_id, estado, fecha_recepcion, creado_en, actualizado_en, talca_id, "
            "cliente_nombre_talca) VALUES (?, 'recibido', '2025-01-15T09:00:00', ?, ?, 8200, 'NOMBRE EN LA OT')",
            (cli, TS, TS)).lastrowid
        for orden, metodo, monto, fecha in [(cls.local, "efectivo", 10000, "2026-11-05"),
                                            (cls.talca, "transferencia", 25000, "2026-11-20"),
                                            (cls.talca, "transferencia", 7000, "2026-12-01")]:
            db.execute("INSERT INTO pagos_orden (orden_id, metodo_pago, monto, fecha, usuario, creado_en) "
                       "VALUES (?, ?, ?, ?, 'cajero', ?)", (orden, metodo, monto, fecha, TS))
        db.commit()
        db.close()

    def _filas(self, **params):
        resp = self.client.get("/api/reportes/cobros", query_string=params)
        if resp.status_code != 200:
            self.fail(f"HTTP {resp.status_code}: {resp.data[:300]!r}")
        filas = list(load_workbook(io.BytesIO(resp.data)).active.iter_rows(values_only=True))
        i = next(k for k, f in enumerate(filas) if f and f[0] == "ID pago")
        return [dict(zip(filas[i], f)) for f in filas[i + 1:] if f and f[0] not in (None, "")]

    def test_filtra_por_fecha_de_pago(self):
        filas = self._filas(desde="2026-11-01", hasta="2026-11-30")
        self.assertEqual([25000, 10000], [f["Monto"] for f in filas])  # más reciente primero

    def test_muestra_el_numero_de_talca_y_el_nombre_de_la_ot(self):
        f = next(x for x in self._filas(desde="2026-11-20", hasta="2026-11-20"))
        self.assertEqual("8200", str(f["N° OT"]))
        self.assertEqual("NOMBRE EN LA OT", f["Cliente"])
        self.assertEqual("transferencia", f["Tipo de cobro"])

    def test_una_ot_local_sale_con_prefijo_L(self):
        f = next(x for x in self._filas(desde="2026-11-05", hasta="2026-11-05"))
        self.assertEqual(f"L-{self.local}", f["N° OT"])

    def test_el_resumen_financiero_ya_no_muestra_el_id_interno(self):
        resp = self.client.get("/api/reportes/finanzas-resumen", query_string={"desde": "2026-11-20", "hasta": "2026-11-20"})
        self.assertEqual(200, resp.status_code)
        ws = load_workbook(io.BytesIO(resp.data))["Ingresos (pagos)"]
        otes = [r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
        self.assertIn("8200", [str(x) for x in otes])
        self.assertNotIn(f"#{self.talca}", otes)

    def test_pdf_responde(self):
        resp = self.client.get("/api/reportes/cobros/pdf")
        self.assertEqual(200, resp.status_code)
        self.assertTrue(resp.data.startswith(b"%PDF"))


if __name__ == "__main__":
    unittest.main()
